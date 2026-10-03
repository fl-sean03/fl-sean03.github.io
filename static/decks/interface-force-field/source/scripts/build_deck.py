#!/usr/bin/env python3
from pathlib import Path
import json,math,csv,sys,os
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION,XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from story import SL,SOURCES
ROOT=Path(__file__).resolve().parents[1];os.chdir(ROOT)
BG='FBFAF7';FG='1B1B1A';MUT='5C5A55';TEAL='166D66';GOLD='80530C';BLUE='315FA2';RED='A23447';CARD='F3F0E8';LINE='B9B3A5';FONT='Lato'

def col(s):return RGBColor.from_string(s)
def box(sl,x,y,w,h,fill=CARD,line=None,rnd=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rnd else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=col(fill);sh.line.fill.background() if line is None else None
 if line:sh.line.color.rgb=col(line)
 if rnd:
  try:sh.adjustments[0]=.08
  except Exception:pass
 return sh

def tx(sl,text,x,y,w,h,size=20,color=FG,bold=False,align=None):
 sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=col(color);p.space_after=Pt(5)
  if align is not None:p.alignment=align
 return sh

def ln(sl,x1,y1,x2,y2,color=MUT,width=1.5):
 sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=col(color);sh.line.width=Pt(width);return sh

def arrow(sl,x,y,w=.45,h=.22,color=TEAL):
 sh=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=col(color);sh.line.fill.background();return sh

def pic(sl,name,x,y,w,h):
 from PIL import Image
 p=ROOT/'media'/name;im=Image.open(p);iw,ih=im.size;s=min(w/iw,h/ih);aw=iw*s;ah=ih*s
 return sl.shapes.add_picture(str(p),Inches(x+(w-aw)/2),Inches(y+(h-ah)/2),width=Inches(aw),height=Inches(ah))

def bullet(sl,text,x,y,w,size=20,num=None):
 box(sl,x,y+.08,.06,.3,TEAL);return tx(sl,text,x+.23,y,w-.23,.92,size)

def base(pr,d,n,total):
 sl=pr.slides.add_slide(pr.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=col(BG)
 tx(sl,'INTERFACE FORCE FIELD',.65,.28,5,.2,10,MUT,bold=True);tx(sl,d['chapter'],8.0,.28,4.65,.24,10,TEAL,align=PP_ALIGN.RIGHT)
 if d['layout']!='hero':
  tx(sl,d['title'],.65,.82,12.0,.75,32,FG,True);tx(sl,d['lead'],.65,1.6,12,.61,19,MUT)
 ln(sl,.65,6.93,12.68,6.93,LINE,.7);box(sl,.65,6.93,12.03*n/total,.025,TEAL)
 src=' · '.join(r for r in d['refs'] if r!='DERIVED') or 'Original illustration'
 tx(sl,src,.65,7.05,10,.2,10,MUT);tx(sl,f'{n:02d} / {total:02d}',11.6,7.02,1.05,.3,11,MUT,align=PP_ALIGN.RIGHT)
 if d.get('tag'):tx(sl,d['tag'],.65,6.48,12,.34,10,GOLD,bold=True)
 return sl

def cards(sl,items,y=2.48,h=3.55):
 n=len(items);gap=.2;w=(12.03-gap*(n-1))/n
 for i,item in enumerate(items):
  x=.65+i*(w+gap);box(sl,x,y,w,h,rnd=True);box(sl,x,y,.06,h,[TEAL,BLUE,GOLD,RED][i%4]);tx(sl,item[0],x+.23,y+.25,w-.46,.8,23,[TEAL,BLUE,GOLD,RED][i%4],True)
  tx(sl,'\n'.join(item[1:]),x+.23,y+1.25,w-.46,h-1.45,21,FG)

def draw(pr,d,n,total):
 sl=base(pr,d,n,total);typ=d['layout'];items=d['items']
 if typ=='hero':
  pic(sl,d['image'],5.55,1.03,7.15,5.25)
  tx(sl,'Interface\nForce Field',.65,1.25,6.9,1.8,50,FG,True);tx(sl,d['lead'],.68,3.42,5.5,1.04,27,TEAL)
  tx(sl,'A visual guide to structure, simulation\nand experimental evidence',.68,4.84,5.5,.85,19,MUT)
  tx(sl,'Sean Florez  ·  3 October 2026',.68,5.97,5.6,.3,13,MUT)
 elif typ=='image':
  for i,t in enumerate(items):bullet(sl,t,.65,2.55+i*1.1,5.0,size=20)
  
  if d['media']:
   sl.shapes.add_movie(str(ROOT/'media'/f'{d["media"]}.mp4'),Inches(5.9),Inches(2.35),Inches(6.8),Inches(3.825),poster_frame_image=str(ROOT/'media'/f'{d["media"]}-poster.png'),mime_type='video/mp4')
  else:pic(sl,d['image'],5.9,2.26,6.8,4.04)
  if d['media']:tx(sl,'Animation is included in the media companion.',6.12,6.18,6.15,.22,11,TEAL)
 elif typ=='chain':
  w=1.88
  for i,(title,sub) in enumerate(items):
   x=.65+i*2.03;box(sl,x,2.9,w,2.1,rnd=True);tx(sl,f'{i+1:02}',x+.19,3.13,w-.38,.5,27,TEAL,True);tx(sl,title,x+.19,3.81,w-.38,.7,21,FG,True);tx(sl,sub,x+.19,4.48,w-.38,.55,13,MUT)
   if i<5:arrow(sl,x+w-.03,3.32,.23,.16)
  tx(sl,'Separate the input, model, calculation and evidence for every claim.',.65,5.69,11.7,.5,22,FG)
 elif typ=='layers':
  for i,(t,b) in enumerate(items):
   y=2.43+i*.84;box(sl,.65,y,12.03,.69,rnd=True);tx(sl,str(i+1),.87,y+.15,.4,.4,23,TEAL,True);tx(sl,t,1.6,y+.15,3.35,.43,22,FG,True);tx(sl,b,5.25,y+.16,7.12,.48,18,MUT)
 elif typ=='family':
  box(sl,.65,2.49,12.03,3.73,rnd=True);tx(sl,items[0][0],.93,2.77,10.6,.48,26,FG,True);tx(sl,items[0][1],.93,3.42,10.9,.7,21,MUT)
  box(sl,1.0,4.44,11.15,1.3,CARD,line=TEAL,rnd=True);tx(sl,'IFF',1.3,4.72,1.7,.55,32,TEAL,True);tx(sl,items[1][1],3.19,4.73,8.67,.7,22,FG)
 elif typ in ('terms','state','io','closing','gates'):
  if typ in ('state','gates'):
   for i,(t,b) in enumerate(items):
    x=.65+(i%2)*6.13;y=2.38+(i//2)*1.82;box(sl,x,y,5.9,1.64,rnd=True);tx(sl,t,x+.23,y+.22,5.44,.48,24,[TEAL,BLUE,GOLD,FG][i],True);tx(sl,b,x+.23,y+.83,5.44,.6,20,FG)
  else:cards(sl,items)
 elif typ=='parameters':
  cards(sl,items,y=2.45,h=2.93);tx(sl,'Rmin governs distance · ε governs well depth',.87,5.7,11.9,.5,25,FG)
 elif typ=='convention':
  for i,(t,b) in enumerate(items):
   y=2.48+i*1.09;tx(sl,t,.7,y,3.5,.45,20,[TEAL,BLUE,GOLD][i],True);tx(sl,b,4.11,y-.02,8.26,.77,26,FG);ln(sl,.65,y+.86,12.65,y+.86,LINE,.8)
 elif typ=='minimize':
  # A deliberately abstract basin, without fake sampled energies.
  for i,t in enumerate(items):bullet(sl,t,.65,2.51+i*1.04,5.2)
  pts=[]
  for i in range(70):
   u=i/69;x=6.38+5.7*u;y=3.03+2.28*(1-(2*u-1)**2)+.2*math.sin(8*u);pts.append((x,y))
  for a,b in zip(pts,pts[1:]):ln(sl,*a,*b,TEAL,2.2)
  for idx in [12,22,31]:
   x,y=pts[idx];sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-.065),Inches(y-.065),Inches(.13),Inches(.13));sh.fill.solid();sh.fill.fore_color.rgb=col(GOLD);sh.line.fill.background()
  tx(sl,'configuration coordinate →',6.5,5.6,5.6,.35,15,MUT);tx(sl,'local basin',9.02,5.43,2.2,.4,19,GOLD)
 elif typ=='movie':
  # Native embedded MP4 with explicit poster, supported separately in browser and static PDF.
  path=ROOT/'media'/f'{d["media"]}.mp4';poster=ROOT/'media'/f'{d["media"]}-poster.png'
  sl.shapes.add_movie(str(path),Inches(5.75),Inches(2.35),Inches(6.89),Inches(3.88),poster_frame_image=str(poster),mime_type='video/mp4')
  for i,t in enumerate(items):bullet(sl,t,.65,2.52+i*1.06,4.84,19)
 elif typ=='observables':
  for i,(t,eq,limit) in enumerate(items):
   x=.65+(i%2)*6.13;y=2.38+(i//2)*1.85;box(sl,x,y,5.9,1.65,rnd=True);tx(sl,t,x+.22,y+.13,5.45,.36,20,TEAL,True);tx(sl,eq,x+.22,y+.63,5.45,.42,25,FG);tx(sl,limit,x+.22,y+1.2,5.45,.3,14,MUT)
 elif typ=='cleavage':
  pic(sl,d['image'],6.24,2.28,6.4,3.17)
  tx(sl,'γ ≈ ΔE / 2A',6.84,5.49,5.1,.5,31,TEAL,True)
  for i,t in enumerate(items):bullet(sl,t,.65,2.57+i*1.07,5.28,20)
 elif typ=='alloy':
  for i,t in enumerate(items):bullet(sl,t,.65,2.54+i*1.04,5.48,20)
  # B2 topology schematic: 8 first-shell Al sites around the central Ni, no metric axes.
  xy=[(7.2,3.1),(9.75,2.47),(11.68,3.26),(9.09,3.95),(7.2,4.75),(9.75,4.15),(11.68,4.91),(9.09,5.62)];center=(9.46,4.13)
  for p in xy:ln(sl,*p,*center,MUT,2)
  for p in xy:
   sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(p[0]-.16),Inches(p[1]-.16),Inches(.32),Inches(.32));sh.fill.solid();sh.fill.fore_color.rgb=col(GOLD);sh.line.fill.background()
  sh=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(center[0]-.23),Inches(center[1]-.23),Inches(.46),Inches(.46));sh.fill.solid();sh.fill.fore_color.rgb=col(BLUE);sh.line.fill.background()
  tx(sl,'Al +0.39e',6.95,5.98,2.7,.29,15,GOLD);tx(sl,'Ni −0.39e',9.95,5.98,2.7,.29,15,BLUE);tx(sl,'B2 topology · schematic · not a metric structure',6.95,6.28,5.6,.25,11,MUT)
 elif typ=='validation':
  for i,(t,b) in enumerate(items[:2]):
   x=.65+i*6.13;c=TEAL if i==0 else GOLD;box(sl,x,2.59,5.9,2.25,rnd=True);tx(sl,t,x+.24,2.88,5.4,.49,26,c,True);tx(sl,b,x+.24,3.76,5.4,.75,22,FG)
  arrow(sl,3.32,5.13,.55,.28,TEAL);arrow(sl,9.48,5.13,.55,.28,GOLD);box(sl,2.2,5.56,8.91,.65,CARD,rnd=True);tx(sl,items[2][1],2.42,5.7,8.43,.36,21,FG,align=PP_ALIGN.CENTER)
 elif typ=='revision':
  for i,t in enumerate(items):bullet(sl,t,.65,2.52+i*1.08,4.95,20)
  sl.shapes.add_movie(str(ROOT/'media'/f'{d["media"]}.mp4'),Inches(5.75),Inches(2.35),Inches(6.89),Inches(3.88),poster_frame_image=str(ROOT/'media'/f'{d["media"]}-poster.png'),mime_type='video/mp4')
 elif typ=='workflow':
  for i,(t,b) in enumerate(items):
   x=.65+(i%4)*3.07;y=2.39+(i//4)*2.0;box(sl,x,y,2.84,1.72,rnd=True);tx(sl,f'{i+1:02}',x+.21,y+.13,.6,.4,17,TEAL,True);tx(sl,t,x+.21,y+.54,2.44,.68,20,FG,True);tx(sl,b,x+.21,y+1.31,2.44,.35,12,MUT)
   if i%4<3:arrow(sl,x+2.84,y+.44,.22,.15)
 elif typ=='status':
  for i,(t,b) in enumerate(items):
   y=2.37+i*1.24;box(sl,.65,y,12.03,1.08,rnd=True);tx(sl,t,.89,y+.19,3.78,.7,21,[TEAL,GOLD,RED][i],True);tx(sl,b,4.95,y+.18,7.38,.76,20,FG)
 elif typ=='comparison':
  w=3.88
  labels=['REPRESENTATION','USEFUL ROLE','WHAT TO CHECK']
  for i,(t,*bs) in enumerate(items):
   x=.65+i*4.07;box(sl,x,2.38,w,3.88,rnd=True);tx(sl,t,x+.24,2.6,w-.48,.5,29,[TEAL,BLUE,GOLD][i],True)
   for j,b in enumerate(bs):tx(sl,labels[j],x+.24,3.37+j*.92,w-.48,.22,10,MUT,True);tx(sl,b,x+.24,3.66+j*.92,w-.48,.57,19,FG)
 elif typ=='loop':
  # Six connected stages, deliberate two-row reading with numbered order.
  for i,(t,b) in enumerate(items):
   x=.65+(i%3)*4.07;y=2.4+(i//3)*1.85;box(sl,x,y,3.83,1.54,rnd=True);tx(sl,f'{i+1}  {t}',x+.23,y+.18,3.37,.55,25,TEAL,True);tx(sl,b,x+.23,y+.9,3.37,.5,18,FG)
   if i%3<2:arrow(sl,x+3.8,y+.55,.3,.18)
 elif typ=='equations':
  for i,(t,eq) in enumerate(items):
   y=2.31+i*.94;tx(sl,t,.7,y,3.73,.42,20,TEAL,True);tx(sl,eq,4.65,y,7.92,.65,23,FG);ln(sl,.65,y+.8,12.65,y+.8,LINE,.7)
 elif typ=='data':
  table=sl.shapes.add_table(4,7,Inches(.65),Inches(2.42),Inches(12.03),Inches(2.22)).table
  headers=['Metal','5a expt / Å','5a 12–6 / Å','γ111 expt / J m⁻²','γ111 12–6','K expt / GPa','K 12–6 / GPa']
  vals=[['Ca (α)','27.942','27.947','0.492 ± 0.01','0.490','20','30'],['Rh','19.016','19.016','2.64 ± 0.02','2.643','276','258'],['Sr (α)','30.420','30.423','0.41 ± 0.01','0.411','12.0','24']]
  for j,h in enumerate(headers):table.cell(0,j).text=h
  for i,row in enumerate(vals,1):
   for j,v in enumerate(row):table.cell(i,j).text=v
  for i,row in enumerate(table.rows):
   for cell in row.cells:
    cell.fill.solid();cell.fill.fore_color.rgb=col(CARD if i else LINE);cell.margin_left=Inches(.1);cell.margin_top=Inches(.09)
    for p in cell.text_frame.paragraphs:p.font.name=FONT;p.font.size=Pt(13);p.font.color.rgb=col(FG);p.font.bold=i==0
  tx(sl,'Not fitted: bulk modulus K / GPa',.65,4.94,5.1,.35,21,TEAL,True)
  # Editable native PowerPoint result chart, backed by embedded spreadsheet.
  cd=CategoryChartData();cd.categories=['Ca','Rh','Sr'];cd.add_series('Experiment',[20,276,12]);cd.add_series('12–6 LJ',[30,258,24]);cd.add_series('9–6 LJ',[21,175,16]);chart=sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(6.0),Inches(4.91),Inches(6.6),Inches(1.51),cd).chart
  chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.RIGHT;chart.legend.font.name=FONT;chart.legend.font.size=Pt(10);chart.legend.font.color.rgb=col(FG)
  chart.chart_style=10
  from pptx.oxml.xmlchemy import OxmlElement
  sp=OxmlElement('c:spPr');fill=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val',BG);fill.append(rgb);sp.append(fill);chart._chartSpace.append(sp)
  sp2=OxmlElement('c:spPr');fill2=OxmlElement('a:solidFill');rgb2=OxmlElement('a:srgbClr');rgb2.set('val',BG);fill2.append(rgb2);sp2.append(fill2);chart._chartSpace.chart.plotArea.append(sp2)
  for i,series in enumerate(chart.series):series.format.fill.solid();series.format.fill.fore_color.rgb=col([FG,TEAL,GOLD][i]);series.format.line.fill.background()
  for ax in [chart.category_axis,chart.value_axis]:ax.tick_labels.font.name=FONT;ax.tick_labels.font.size=Pt(10);ax.tick_labels.font.color.rgb=col(MUT)
  chart.value_axis.minimum_scale=0
  chart.value_axis.maximum_scale=320
  tx(sl,'K references are selected, not averaged.\nSI S7-S8: repeatability about ±3%.\nProtocol agreement is separate from accuracy.',.65,5.46,5.08,.90,14,MUT)
 elif typ=='references':
  labels={'K21':'Kanhaiya et al. (2021) · FCC-metal parameters, methods and results','K21C':'Author Correction (2021) · restored structure and script data','L18':'Liu et al. (2018) · alloy bonding (primary paper citation)','L18SI':'Liu Supporting Information · S15–S19, Tables S1–S2','LJ':'LAMMPS · lj/cut parameter convention','MD':'LAMMPS · velocity Verlet / minimization','MACE':'Batatia et al. (2022) · MACE (retrieved v2, 2023)','AGENT':'IFF Agent workflow documentation (2026)'}
  for i,r in enumerate(d['refs']):
   src=SOURCES[r];y=2.32+i*.47;sh=tx(sl,labels[r],.7,y,11.94,.4,17,FG)
   if src['url']:
    sh.click_action.hyperlink.address=src['url'];sh.text_frame.paragraphs[0].font.color.rgb=col(TEAL)
 else:raise ValueError(typ)
 note=d['notes']+'\n\nEVIDENCE / SOURCES\n'+ '\n'.join(f'{r}: {SOURCES[r]["title"]}\n{SOURCES[r]["url"] or "Implemented workflow documentation / original derivation"}\nScope: {SOURCES[r]["scope"]}' for r in d['refs'])
 if d['media']:note+=f'\n\nMedia companion: {d["media"]}.mp4, .gif and -poster.png. Static visuals remain authoritative for interpretation.'
 sl.notes_slide.notes_text_frame.text=note
 return sl

def rewrite_application_metadata(dest,slides):
 """Replace template application claims with observed authoring facts."""
 import zipfile
 from lxml import etree
 ns='http://schemas.openxmlformats.org/officeDocument/2006/extended-properties'
 vt='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
 root=etree.Element('{'+ns+'}Properties',nsmap={None:ns,'vt':vt})
 values={'PresentationFormat':'Widescreen (16:9)','Slides':str(len(slides)),'Notes':str(len(slides)),'HiddenSlides':'0','MMClips':str(sum(bool(d.get('media')) for d in slides)),'Company':'','Manager':''}
 for k,v in values.items():etree.SubElement(root,'{'+ns+'}'+k).text=v
 raw=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(dest) as z:parts=[(i,z.read(i.filename)) for i in z.infolist()]
 with zipfile.ZipFile(dest,'w') as z:
  for i,data in parts:z.writestr(i,raw if i.filename=='docProps/app.xml' else data)

def build():
 sel=SL
 pr=Presentation();pr.slide_width=Inches(13.333333);pr.slide_height=Inches(7.5)
 from datetime import datetime
 metadata=json.loads((ROOT/'scripts/production_metadata.json').read_text())
 cp=pr.core_properties
 cp.title=metadata['title'];cp.author=metadata['author'];cp.last_modified_by=''
 cp.created=cp.modified=datetime.fromisoformat(metadata['timestamp_utc'])
 cp.comments=metadata['description'];cp.subject='Interface Force Field and the implemented evidence workflow'
 cp.keywords='IFF, molecular mechanics, force field, experimental calibration, validation';cp.revision=2
 cp.category='Educational scientific presentation'

 for i,d in enumerate(sel,1):draw(pr,d,i,len(sel))
 (ROOT/'build').mkdir(exist_ok=True)
 dest=ROOT/'build'/'iff-visual-showcase.pptx';pr.save(dest)
 rewrite_application_metadata(dest,sel)
 notes=[]
 for i,d in enumerate(sel,1):notes.append(f'## {i:02d}. {d["title"]}\n\n{d["notes"]}\n\n'+ '\n'.join(f'- [{SOURCES[r]["title"]}]({SOURCES[r]["url"]})' if SOURCES[r]['url'] else '- '+SOURCES[r]['title'] for r in d['refs']))
 (dest.with_suffix('.notes.md')).write_text('# Presenter notes and citations\n\n'+'\n\n'.join(notes))
 print(dest, len(pr.slides),'slides')
 data=[dict(d,deck_slide=i+1) for i,d in enumerate(sel)]
 (dest.with_suffix('.slides.json')).write_text(json.dumps(data,indent=2,ensure_ascii=False))
 return dest
if __name__=='__main__':
 build()
